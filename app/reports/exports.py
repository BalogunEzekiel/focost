from io import BytesIO
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_RIGHT
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak


class ExportService:
    @staticmethod
    def _money(value):
        return f"NGN {float(value or 0):,.2f}"

    @staticmethod
    def excel(data, start_date=None, end_date=None, category=None):
        wb = Workbook()
        ws = wb.active
        ws.title = "Report Summary"

        ws["A1"] = "FOCOST — FINANCIAL REPORT"
        ws["A1"].font = Font(size=18, bold=True)
        ws.merge_cells("A1:F1")
        ws["A2"] = "Generated"
        ws["B2"] = datetime.now().strftime("%d %b %Y %H:%M")
        ws["A3"] = "Period"
        ws["B3"] = f"{start_date or 'All time'} to {end_date or 'Present'}"
        ws["A4"] = "Category"
        ws["B4"] = category or "All Categories"

        rows = [
            ("Total Income", data["summary"]["income"]),
            ("Total Expenses", data["summary"]["expenses"]),
            ("Net Savings", data["summary"]["savings"]),
            ("Savings Rate (%)", data["summary"]["savings_rate"]),
            ("Transactions", data["summary"]["transactions"]),
        ]
        for r, (label, value) in enumerate(rows, 6):
            ws.cell(r, 1, label).font = Font(bold=True)
            ws.cell(r, 2, value)
            if r != 9 and r != 10:
                ws.cell(r, 2).number_format = '#,##0.00'

        for col, width in {1: 28, 2: 20, 3: 18, 4: 18, 5: 18, 6: 18}.items():
            ws.column_dimensions[get_column_letter(col)].width = width

        # Category sheets
        for title, key in (("Income by Category", "income_categories"), ("Expense by Category", "expense_categories")):
            sh = wb.create_sheet(title[:31])
            sh.append(["Category", "Amount"])
            for cell in sh[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
            for label, value in zip(data[key]["labels"], data[key]["values"]):
                sh.append([label, value])
            sh.column_dimensions["A"].width = 32
            sh.column_dimensions["B"].width = 20
            for row in sh.iter_rows(min_row=2, min_col=2, max_col=2):
                row[0].number_format = '#,##0.00'

        # Transactions
        sh = wb.create_sheet("Transactions")
        sh.append(["Type", "Date", "Category", "Source / Merchant", "Amount", "Recurring", "Notes"])
        for cell in sh[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for item in data["summary"]["income_records"]:
            sh.append(["Income", item.received_date, item.category, item.source, float(item.amount or 0), "Yes" if item.recurring else "No", item.notes or ""])
        for item in data["summary"]["expense_records"]:
            sh.append(["Expense", item.expense_date, item.category, item.merchant, float(item.amount or 0), "Yes" if item.recurring else "No", item.notes or ""])
        widths = [12, 15, 24, 30, 18, 12, 40]
        for i, width in enumerate(widths, 1):
            sh.column_dimensions[get_column_letter(i)].width = width
        for row in sh.iter_rows(min_row=2):
            row[1].number_format = "dd mmm yyyy"
            row[4].number_format = '#,##0.00'

        out = BytesIO()
        wb.save(out)
        out.seek(0)
        return out

    @staticmethod
    def pdf(data, start_date=None, end_date=None, category=None):
        out = BytesIO()
        doc = SimpleDocTemplate(out, pagesize=landscape(A4), rightMargin=28, leftMargin=28, topMargin=28, bottomMargin=28)
        styles = getSampleStyleSheet()
        styles.add(ParagraphStyle(name="RightSmall", parent=styles["Normal"], alignment=TA_RIGHT, fontSize=8))
        story = [
            Paragraph("FOCOST — FINANCIAL REPORT", styles["Title"]),
            Paragraph(f"Generated: {datetime.now().strftime('%d %b %Y %H:%M')} &nbsp; | &nbsp; Period: {start_date or 'All time'} to {end_date or 'Present'} &nbsp; | &nbsp; Category: {category or 'All Categories'}", styles["Normal"]),
            Spacer(1, 14),
        ]
        summary = data["summary"]
        kpis = [["Total Income", "Total Expenses", "Net Savings", "Savings Rate", "Transactions"],
                [ExportService._money(summary["income"]), ExportService._money(summary["expenses"]), ExportService._money(summary["savings"]), f"{summary['savings_rate']:.1f}%", str(summary["transactions"])]]
        t = Table(kpis, colWidths=[145] * 5)
        t.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#1f4e78")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"), ("ALIGN", (0,0), (-1,-1), "CENTER"), ("BOX", (0,0), (-1,-1), .5, colors.grey), ("INNERGRID", (0,0), (-1,-1), .25, colors.lightgrey), ("TOPPADDING", (0,0), (-1,-1), 8), ("BOTTOMPADDING", (0,0), (-1,-1), 8)]))
        story += [t, Spacer(1, 18)]

        for title, key in (("Income by Category", "income_categories"), ("Expense by Category", "expense_categories")):
            table_data = [["Category", "Amount"]] + [[str(a), ExportService._money(b)] for a, b in zip(data[key]["labels"], data[key]["values"])]
            if len(table_data) == 1:
                table_data.append(["No data", "NGN 0.00"])
            table = Table(table_data, colWidths=[240, 130], repeatRows=1)
            table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#1f4e78")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"), ("ALIGN", (1,1), (1,-1), "RIGHT"), ("GRID", (0,0), (-1,-1), .25, colors.lightgrey), ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#f6f8fa")]), ("TOPPADDING", (0,0), (-1,-1), 6), ("BOTTOMPADDING", (0,0), (-1,-1), 6)]))
            story += [Paragraph(title, styles["Heading2"]), table, Spacer(1, 14)]

        story.append(PageBreak())
        story.append(Paragraph("Transaction Detail", styles["Heading1"]))
        tx = [["Type", "Date", "Category", "Source / Merchant", "Amount"]]
        for item in summary["income_records"]:
            tx.append(["Income", item.received_date.strftime("%d %b %Y"), str(item.category or "Uncategorized"), str(item.source or ""), ExportService._money(item.amount)])
        for item in summary["expense_records"]:
            tx.append(["Expense", item.expense_date.strftime("%d %b %Y"), str(item.category or "Uncategorized"), str(item.merchant or ""), ExportService._money(item.amount)])
        if len(tx) == 1:
            tx.append(["—", "—", "No transactions", "—", "NGN 0.00"])
        table = Table(tx, colWidths=[65, 90, 130, 250, 100], repeatRows=1)
        table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#1f4e78")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"), ("GRID", (0,0), (-1,-1), .25, colors.lightgrey), ("ALIGN", (-1,1), (-1,-1), "RIGHT"), ("FONTSIZE", (0,0), (-1,-1), 8), ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#f6f8fa")])]))
        story.append(table)
        doc.build(story)
        out.seek(0)
        return out
