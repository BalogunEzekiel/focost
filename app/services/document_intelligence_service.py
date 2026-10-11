"""Private document intake, extraction, staging, deduplication and confirmed import."""

import csv
import hashlib
import io
import json
import os
import re
import statistics
import uuid
from datetime import date, datetime
from pathlib import Path

from flask import current_app
from werkzeug.utils import secure_filename
from sqlalchemy.exc import IntegrityError

from app.extensions import db
from app.models.integration import DocumentImport
from app.models.expense import Expense
from app.models.income import Income
from app.services.category_service import CategoryService


ALLOWED = {
    "csv", "xlsx", "xlsm", "pdf",
    "png", "jpg", "jpeg", "webp", "tif", "tiff",
}
MAX_BYTES = 10 * 1024 * 1024
MAX_PDF_PAGES = 50
MAX_OCR_PAGES = 10
MAX_TRANSACTIONS = 1000

DEFAULT_TESSERACT_CMD = (
    r"C:\Program Files\Tesseract-OCR\tesseract.exe"
)


def _date(value):
    """Normalize supported date values to ISO format."""
    if isinstance(value, datetime):
        return value.date().isoformat()

    if isinstance(value, date):
        return value.isoformat()

    value = str(value or "").strip()

    if not value:
        return None

    formats = (
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%m/%d/%Y",
        "%d-%m-%Y",
        "%m-%d-%Y",
        "%Y/%m/%d",
        "%d %b %Y",
        "%d %B %Y",
        "%d/%m/%y",
        "%m/%d/%y",
        "%d-%m-%y",
        "%d %b, %Y",
        "%d-%b-%Y",
        "%d-%B-%Y",
        "%d-%b-%y",
        "%d-%B-%y",
    )

    for fmt in formats:
        try:
            return datetime.strptime(value[:40], fmt).date().isoformat()
        except ValueError:
            continue

    match = re.search(
        r"\b(20\d{2})[-/](\d{1,2})[-/](\d{1,2})\b",
        value,
    )

    if match:
        try:
            return date(
                int(match.group(1)),
                int(match.group(2)),
                int(match.group(3)),
            ).isoformat()
        except ValueError:
            pass

    return None


def _amount(value):
    """Convert a monetary value into a signed float."""
    if value is None:
        return None

    value = str(value).strip()

    if not value:
        return None

    negative_parentheses = (
        value.startswith("(") and value.endswith(")")
    )

    cleaned = (
        value.replace(",", "")
        .replace("₦", "")
        .replace("$", "")
        .replace("NGN", "")
        .replace("ngn", "")
        .replace(" ", "")
    )

    if negative_parentheses:
        cleaned = "-" + cleaned[1:-1]

    # Permit only digits, decimal points and minus signs.
    cleaned = re.sub(r"[^0-9.\-]", "", cleaned)

    try:
        number = round(float(cleaned), 2)
        return number if number != 0 else None
    except (ValueError, TypeError, OverflowError):
        return None


def _category(text, kind):
    """Suggest a category without automatically confirming the transaction."""
    text = (text or "").lower()

    if kind == "income":
        if any(word in text for word in ("salary", "payroll", "wage")):
            return "Salary"

        if any(word in text for word in ("dividend", "interest")):
            return "Investment"

        if any(word in text for word in ("refund", "reversal")):
            return "Other"

        return "Other"

    mapping = [
        ("Food", (
            "food", "restaurant", "grocery", "market",
            "supermarket", "cafe",
        )),
        ("Transportation", (
            "uber", "bolt", "bus", "transport", "taxi", "flight",
        )),
        ("Housing", ("rent", "landlord", "housing")),
        ("Utilities", ("electric", "power", "water", "utility")),
        ("Telephone", ("airtime", "data", "telecom", "mtn", "airtel")),
        ("Healthcare", ("pharmacy", "hospital", "clinic", "medical")),
        ("Education", ("school", "tuition", "course", "book")),
        ("Petrol", ("fuel", "petrol", "filling station")),
        ("Shopping", ("shop", "store", "purchase", "pos")),
    ]

    return next(
        (
            category
            for category, words in mapping
            if any(word in text for word in words)
        ),
        "Other",
    )


def _fingerprint(day, amount, description):
    """Create a stable fingerprint for possible duplicate detection."""
    normalized_description = " ".join(
        (description or "").lower().split()
    )

    source = (
        f"{day}|{float(amount):.2f}|{normalized_description}"
    )

    return hashlib.sha256(source.encode("utf-8")).hexdigest()


class DocumentIntelligenceService:
    """Extract, stage, review and import financial documents."""

    @staticmethod
    def _configure_tesseract():
        """Configure the Tesseract executable for Windows or other hosts."""
        import pytesseract

        configured_path = (
            current_app.config.get("TESSERACT_CMD")
            or os.environ.get("TESSERACT_CMD")
            or DEFAULT_TESSERACT_CMD
        )

        configured_path = str(configured_path).strip()

        if configured_path and not Path(configured_path).is_file():
            # Allow the executable to be found through PATH if a bare
            # executable name was explicitly configured.
            if (
                os.path.basename(configured_path).lower()
                == configured_path.lower()
            ):
                pytesseract.pytesseract.tesseract_cmd = configured_path
                return pytesseract

            raise ValueError(
                "Tesseract OCR was not found at the configured path. "
                "Check TESSERACT_CMD in your FOCOST .env file."
            )

        if configured_path:
            pytesseract.pytesseract.tesseract_cmd = configured_path

        return pytesseract

    @staticmethod
    def _extract_pdf_ocr(raw):
        """OCR scanned or mixed PDFs using PyMuPDF and Tesseract."""
        try:
            import pymupdf
            from PIL import Image
        except ImportError as exc:
            raise ValueError(
                "Scanned PDF extraction requires PyMuPDF, Pillow, "
                "pytesseract and the Tesseract OCR engine."
            ) from exc

        pytesseract = (
            DocumentIntelligenceService._configure_tesseract()
        )

        try:
            with pymupdf.open(stream=raw, filetype="pdf") as doc:
                if len(doc) > MAX_OCR_PAGES:
                    raise ValueError(
                        f"Scanned PDF OCR is limited to {MAX_OCR_PAGES} "
                        "pages per upload. Split the document and upload "
                        "each part separately."
                    )

                chunks = []

                for page in doc:
                    pixmap = page.get_pixmap(
                        matrix=pymupdf.Matrix(2, 2),
                        alpha=False,
                    )

                    image = Image.open(
                        io.BytesIO(pixmap.tobytes("png"))
                    ).convert("RGB")

                    page_text = pytesseract.image_to_string(
                        image,
                        lang="eng",
                        timeout=30,
                    )

                    chunks.append(page_text)

                return "\n".join(chunks)

        except ValueError:
            raise
        except Exception as exc:
            current_app.logger.exception(
                "PDF OCR extraction failed"
            )
            raise ValueError(
                "PDF OCR failed. Verify that Tesseract OCR and the "
                "English language data are installed and configured, "
                "and that the PDF pages are readable."
            ) from exc

    @staticmethod
    def _extract_text(ext, raw):
        """Extract embedded text from supported documents."""
        if ext == "csv":
            try:
                return raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                return raw.decode("latin-1")

        if ext in {"xlsx", "xlsm"}:
            try:
                from openpyxl import load_workbook

                workbook = load_workbook(
                    io.BytesIO(raw),
                    read_only=True,
                    data_only=True,
                )

                try:
                    worksheet = workbook.active
                    output = io.StringIO()
                    writer = csv.writer(output)

                    for row in worksheet.iter_rows(values_only=True):
                        writer.writerow([
                            str(value) if value is not None else ""
                            for value in row
                        ])

                    return output.getvalue()
                finally:
                    workbook.close()

            except Exception as exc:
                raise ValueError(
                    "The spreadsheet could not be read. "
                    "Please upload a valid .xlsx or .xlsm file."
                ) from exc

        if ext == "pdf":
            try:
                from pypdf import PdfReader

                reader = PdfReader(io.BytesIO(raw))

                if len(reader.pages) > MAX_PDF_PAGES:
                    raise ValueError(
                        f"PDF statements are limited to "
                        f"{MAX_PDF_PAGES} pages per upload. "
                        "Split the document and upload each part separately."
                    )

                # Embedded-text extraction only. OCR fallback is managed
                # by create() to avoid running OCR twice unnecessarily.
                return "\n".join(
                    (page.extract_text() or "")
                    for page in reader.pages[:MAX_PDF_PAGES]
                )

            except ImportError as exc:
                raise ValueError(
                    "PDF extraction requires pypdf. "
                    "Install it and retry."
                ) from exc
            except ValueError:
                raise
            except Exception as exc:
                raise ValueError(
                    "The PDF could not be read. Try exporting it again "
                    "or upload a clear image."
                ) from exc

        if ext in {"png", "jpg", "jpeg", "webp", "tif", "tiff"}:
            try:
                from PIL import Image

                pytesseract = (
                    DocumentIntelligenceService._configure_tesseract()
                )

                image = Image.open(
                    io.BytesIO(raw)
                ).convert("RGB")

                return pytesseract.image_to_string(
                    image,
                    lang="eng",
                    timeout=30,
                )

            except ImportError as exc:
                raise ValueError(
                    "Image OCR requires Pillow, pytesseract and "
                    "the Tesseract OCR engine."
                ) from exc
            except ValueError:
                raise
            except Exception as exc:
                current_app.logger.exception(
                    "Image OCR extraction failed"
                )
                raise ValueError(
                    "OCR could not read this image. Try a sharper, "
                    "well-lit image with all receipt edges visible."
                ) from exc

        raise ValueError("Unsupported file type.")

    @staticmethod
    def _opay_statement_rows(text):
        """Parse OPay account statements without confusing references for money.

        OPay statements contain separate Wallet and Savings sections. Many
        OWealth withdrawal/auto-save entries are the two sides of movements
        between those sections, not new income or spending. These internal
        movements are intentionally omitted from the personal cash-flow rows.
        """
        rows = []
        date_row = re.compile(
            r"(?P<date>\d{1,2}\s+[A-Za-z]{3}\s+20\d{2})"
            r"\s+(?P<time>\d{2}:\d{2}:\d{2})"
            r"\s+(?P<value_date>\d{1,2}\s+[A-Za-z]{3}\s+20\d{2})"
        )
        money = r"(?:\d{1,3}(?:,\d{3})+|\d+)\.\d{2}"
        columns = re.compile(
            rf"(?P<debit>--|{money})\s+"
            rf"(?P<credit>--|{money})\s+"
            rf"(?P<balance>{money})(?=\s|$)",
            re.IGNORECASE,
        )
        matches = list(date_row.finditer(text))

        for index, start in enumerate(matches):
            end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
            block = text[start.end():end]
            column_match = columns.search(block)
            if not column_match:
                # Reject the row rather than guessing from arbitrary digits.
                continue

            description = " ".join(block[:column_match.start()].split())
            description = re.sub(r"\s+", " ", description).strip(" |\t\r\n")
            if not description or len(description) < 2:
                continue

            # OWealth Withdrawal and Auto-save entries are internal transfers
            # represented on both account sections. Counting them as income or
            # expense would inflate cash flow and double-count the movement.
            normalized = description.casefold()
            if "owealth withdrawal" in normalized or "auto-save to owealth balance" in normalized:
                continue

            debit_text = column_match.group("debit")
            credit_text = column_match.group("credit")
            debit = _amount(debit_text) if debit_text != "--" else None
            credit = _amount(credit_text) if credit_text != "--" else None

            # A row must have exactly one populated debit/credit column.
            if (debit is None) == (credit is None):
                continue

            amount = -abs(debit) if debit is not None else abs(credit)
            if not amount or abs(amount) > 100_000_000:
                continue

            transaction_date = _date(start.group("date"))
            if not transaction_date:
                continue

            kind = "income" if amount > 0 else "expense"
            category = _category(description, kind)
            needs_review = False
            review_reason = None

            if "transfer from" in normalized:
                # Transfers from another account/person should not be silently
                # assumed to be earned income.
                category = "Transfer"
                needs_review = True
                review_reason = "Confirm whether this incoming transfer is income or a transfer between your own accounts."
            elif "transfer to" in normalized:
                category = "Transfer"
            elif "stamp duty" in normalized:
                category = "Bank Charges"

            row = {
                "date": transaction_date,
                "description": description[:255],
                "amount": round(amount, 2),
                "category": category,
                "kind": kind,
                "confirmed": False,
                "needs_review": needs_review,
            }
            if review_reason:
                row["review_reason"] = review_reason
            rows.append(row)

            if len(rows) > MAX_TRANSACTIONS:
                break

        return rows

    @staticmethod
    def _statement_rows(text, ext):
        """Parse transaction candidates for user review."""
        # Handle OPay's two-account statement layout before the generic parser.
        # Some PDF extractors join the preceding date directly to the account
        # heading (for example, "30 Sep 2026Wallet Account"), so do not require
        # a word boundary immediately before the heading.
        if (
            ext == "pdf"
            and re.search(r"Wallet\s+Account", text, re.I)
            and re.search(r"Savings\s+Account", text, re.I)
        ):
            rows = DocumentIntelligenceService._opay_statement_rows(text)
            if rows:
                # Preserve the same review metadata used by the existing parser.
                seen = set()
                for row in rows:
                    fingerprint = _fingerprint(
                        row["date"], row["amount"], row["description"]
                    )
                    row["fingerprint"] = fingerprint
                    row["possible_duplicate"] = fingerprint in seen
                    row["reconciliation_warning"] = False
                    row["unusual_activity"] = False
                    seen.add(fingerprint)

                magnitudes = [abs(float(row["amount"])) for row in rows]
                if len(magnitudes) >= 5:
                    median_amount = statistics.median(magnitudes)
                    deviations = [
                        abs(value - median_amount) for value in magnitudes
                    ]
                    mad = statistics.median(deviations)
                    threshold = median_amount + (
                        6 * mad if mad else max(median_amount * 5, 1.0)
                    )
                    for row in rows:
                        row["unusual_activity"] = (
                            abs(float(row["amount"])) > threshold
                        )
                return rows

        rows = []

        if ext in {"csv", "xlsx", "xlsm"}:
            parsed = list(csv.reader(io.StringIO(text)))

            if not parsed:
                return rows

            headers = [
                re.sub(r"[^a-z]", "", str(value).lower())
                for value in parsed[0]
            ]

            def idx(*names):
                return next(
                    (
                        headers.index(name)
                        for name in names
                        if name in headers
                    ),
                    None,
                )

            date_index = idx(
                "date",
                "transactiondate",
                "valuedate",
                "postingdate",
            )

            description_index = idx(
                "description",
                "narration",
                "details",
                "transactiondetails",
                "particulars",
                "merchant",
            )

            amount_index = idx("amount", "transactionamount")

            debit_index = idx(
                "debit", "withdrawal", "moneyout"
            )

            credit_index = idx(
                "credit", "deposit", "moneyin"
            )

            balance_index = idx(
                "balance",
                "runningbalance",
                "closingbalance",
                "availablebalance",
            )

            if (
                date_index is None
                or description_index is None
                or (
                    amount_index is None
                    and debit_index is None
                    and credit_index is None
                )
            ):
                raise ValueError(
                    "Could not identify date, description and "
                    "amount/debit/credit columns. Use a header row "
                    "with those labels."
                )

            for row in parsed[1:]:
                required_indexes = [
                    index
                    for index in (
                        date_index,
                        description_index,
                        amount_index,
                        debit_index,
                        credit_index,
                        balance_index,
                    )
                    if index is not None
                ]

                if not row or max(required_indexes) >= len(row):
                    continue

                transaction_date = _date(row[date_index])
                description = str(
                    row[description_index] or ""
                ).strip()

                if not transaction_date or not description:
                    continue

                amount = (
                    _amount(row[amount_index])
                    if amount_index is not None
                    else None
                )

                if amount is None and debit_index is not None:
                    debit_value = _amount(row[debit_index])

                    if debit_value is not None:
                        amount = -abs(debit_value)

                if amount is None and credit_index is not None:
                    credit_value = _amount(row[credit_index])

                    if credit_value is not None:
                        amount = abs(credit_value)

                if amount is None or amount == 0:
                    continue

                balance = (
                    _amount(row[balance_index])
                    if balance_index is not None
                    else None
                )

                kind = "income" if amount > 0 else "expense"

                rows.append({
                    "date": transaction_date,
                    "description": description[:255],
                    "amount": round(amount, 2),
                    "balance": balance,
                    "category": _category(description, kind),
                    "kind": kind,
                    "confirmed": False,
                })

                if len(rows) > MAX_TRANSACTIONS:
                    break

        else:
            # Bank PDFs often extract each table cell on its own line:
            # date, value date, reference, debit, credit, balance, branch,
            # followed by a description that may wrap over several lines.
            # Parse that column-oriented layout first, retaining blank lines
            # because they identify empty debit/credit cells.
            raw_lines = text.splitlines()
            normalized_lines = [line.strip() for line in raw_lines]

            bank_date_pattern = re.compile(
                r"^(?:\d{1,2}[-/]\d{1,2}[-/]\d{2,4}"
                r"|\d{4}[-/]\d{1,2}[-/]\d{1,2}"
                r"|\d{1,2}[- ](?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)"
                r"[a-z]*[- ,]+\d{4})$",
                re.IGNORECASE,
            )

            transaction_starts = []
            for line_index, line in enumerate(normalized_lines):
                if not bank_date_pattern.fullmatch(line):
                    continue
                next_nonempty = next(
                    (
                        normalized_lines[j]
                        for j in range(line_index + 1, min(line_index + 4, len(raw_lines)))
                        if normalized_lines[j]
                    ),
                    "",
                )
                # A real transaction usually begins with transaction date
                # followed by value date. This excludes statement print dates.
                if _date(line) and _date(next_nonempty):
                    transaction_starts.append(line_index)

            def clean_description(parts):
                ignored_patterns = (
                    r"^this is a computer generated email\b",
                    r"^statement period\b",
                    r"^print\.?\s*date\b",
                    r"^branch name\b",
                    r"^account no\b",
                    r"^internal reference\b",
                    r"^customer statement\b",
                    r"^trans\.?\s*date\b",
                    r"^value\.?\s*date\b",
                    r"^reference\b",
                    r"^debits?\b",
                    r"^credits?\b",
                    r"^balance\b",
                    r"^originating branch\b",
                    r"^remarks\b",
                    r"^fax\b",
                    r"^phone\b",
                )
                kept = []
                for part in parts:
                    part = " ".join(str(part).split()).strip()
                    if not part:
                        continue
                    if any(re.search(pattern, part, re.IGNORECASE) for pattern in ignored_patterns):
                        continue
                    if re.fullmatch(r"\d+\.", part):
                        continue
                    kept.append(part)
                return " ".join(kept).strip()

            for start_pos, line_index in enumerate(transaction_starts):
                block_end = (
                    transaction_starts[start_pos + 1]
                    if start_pos + 1 < len(transaction_starts)
                    else len(raw_lines)
                )
                block = raw_lines[line_index:block_end]

                if len(block) < 7:
                    continue

                transaction_date = _date(block[0].strip())
                value_date = _date(block[1].strip()) if len(block) > 1 else None
                if not transaction_date:
                    continue

                # Locate the branch column, which separates the numeric
                # debit/credit/balance cells from the wrapped narration.
                branch_index = next(
                    (
                        j for j, value in enumerate(block[2:], start=2)
                        if re.match(r"^\s*\d{2,}\s+[A-Za-z]", value)
                    ),
                    None,
                )
                if branch_index is None:
                    # Some statements use a textual branch name without a
                    # numeric branch code (for example, "THE PLAZA").
                    # In the standard layout the branch still follows the
                    # debit, credit and balance cells at offset six.
                    branch_index = 6 if len(block) > 6 else len(block)

                reference_parts = block[2:branch_index]
                numeric_cells = [
                    (j, _amount(value))
                    for j, value in enumerate(reference_parts, start=2)
                    if _amount(value) is not None
                ]
                numeric_cells = [
                    item for item in numeric_cells if item[1] is not None
                ]

                amount = None
                balance = None
                if len(numeric_cells) >= 2:
                    amount_pos, amount_value = numeric_cells[-2]
                    balance = numeric_cells[-1][1]
                    # A blank immediately before the amount denotes an
                    # empty debit cell followed by a credit in columnar
                    # statements such as GTBank's.
                    previous_cell = (
                        block[amount_pos - 1].strip()
                        if amount_pos > 2 else ""
                    )
                    amount = (
                        abs(amount_value)
                        if not previous_cell
                        else -abs(amount_value)
                    )
                elif numeric_cells:
                    amount = numeric_cells[-1][1]
                    previous_cell = (
                        block[numeric_cells[-1][0] - 1].strip()
                        if numeric_cells[-1][0] > 2 else ""
                    )
                    amount = abs(amount) if not previous_cell else -abs(amount)

                branch = (
                    block[branch_index].strip()
                    if branch_index < len(block) else ""
                )
                description_parts = (
                    block[branch_index + 1:]
                    if branch_index < len(block) else []
                )
                description = clean_description(description_parts)
                reference = " ".join(
                    part.strip() for part in block[2:branch_index]
                    if part.strip() and _amount(part) is None
                )

                if (
                    not transaction_date
                    or not description
                    or amount is None
                    or amount == 0
                ):
                    continue

                # Do not treat statement totals or balance summaries as
                # individual transactions.
                if re.search(
                    r"^(opening balance|closing balance|total debit|"
                    r"total credit|grand total|available balance)\b",
                    description,
                    re.IGNORECASE,
                ):
                    continue

                kind = "income" if amount > 0 else "expense"
                row = {
                    "date": transaction_date,
                    "description": description[:255],
                    "amount": round(amount, 2),
                    "balance": balance,
                    "category": _category(description, kind),
                    "kind": kind,
                    "confirmed": False,
                }

                # Credits labelled reversal/refund are not necessarily
                # earned income. Preserve them for explicit user review.
                if re.search(r"\b(reversed|reversal|refund|reversal of)\b", description, re.IGNORECASE):
                    row["needs_review"] = True
                    row["review_note"] = (
                        "This appears to be a reversal or refund. "
                        "Verify its accounting treatment before confirming."
                    )

                rows.append(row)
                if len(rows) > MAX_TRANSACTIONS:
                    break

            # Fallback for OCR and simpler PDFs that do not expose a
            # column-oriented table. Keep the legacy date-first parser.
            if not rows:
                date_pattern = (
                    r"(?:"
                    r"\d{1,2}[/-]\d{1,2}[/-]\d{2,4}"
                    r"|20\d{2}-\d{1,2}-\d{1,2}"
                    r"|\d{1,2}[- ](?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*[- ,]+\d{4}"
                    r")"
                )
                amount_pattern = (
                    r"(?P<amount>\(?-?[₦$]?\s*[\d,]+(?:\.\d{1,2})?\)?)\s*$"
                )
                line_pattern = re.compile(
                    rf"^\s*(?P<date>{date_pattern})"
                    rf"\s+(?P<description>.{{2,200}}?)"
                    rf"\s+{amount_pattern}",
                    re.IGNORECASE,
                )

                for original_line in raw_lines:
                    line = " ".join(original_line.split())
                    if not line:
                        continue
                    match = line_pattern.search(line)
                    if not match:
                        continue
                    transaction_date = _date(match.group("date"))
                    description = match.group("description").strip()
                    amount = _amount(match.group("amount"))
                    if (
                        not transaction_date
                        or not description
                        or amount is None
                        or amount == 0
                    ):
                        continue
                    if re.search(
                        r"^(date|transaction date|opening balance|"
                        r"closing balance|total|grand total)\b",
                        description,
                        re.IGNORECASE,
                    ):
                        continue
                    kind = "income" if amount > 0 else "expense"
                    rows.append({
                        "date": transaction_date,
                        "description": description[:255],
                        "amount": round(amount, 2),
                        "category": _category(description, kind),
                        "kind": kind,
                        "confirmed": False,
                    })
                    if len(rows) > MAX_TRANSACTIONS:
                        break

        # Duplicate flags are advisory; rows are not silently discarded.
        seen = set()

        for row in rows:
            fingerprint = _fingerprint(
                row["date"],
                row["amount"],
                row["description"],
            )

            row["fingerprint"] = fingerprint
            row["possible_duplicate"] = fingerprint in seen
            row["reconciliation_warning"] = False
            row["unusual_activity"] = False

            seen.add(fingerprint)

        # Robust outlier flag; this does not reject transactions.
        magnitudes = [
            abs(float(row["amount"]))
            for row in rows
        ]

        if len(magnitudes) >= 5:
            median_amount = statistics.median(magnitudes)

            deviations = [
                abs(value - median_amount)
                for value in magnitudes
            ]

            mad = statistics.median(deviations)

            threshold = median_amount + (
                6 * mad
                if mad
                else max(median_amount * 5, 1.0)
            )

            for row in rows:
                row["unusual_activity"] = (
                    abs(float(row["amount"])) > threshold
                )

        # Reconcile consecutive balances when balances are available.
        ordered = sorted(
            (
                row for row in rows
                if row.get("balance") is not None
            ),
            key=lambda row: row["date"],
        )

        for previous, current in zip(ordered, ordered[1:]):
            expected_delta = float(current["amount"])

            observed_delta = (
                float(current["balance"])
                - float(previous["balance"])
            )

            if abs(expected_delta - observed_delta) > 0.02:
                current["reconciliation_warning"] = True

        return rows

        
    @staticmethod
    def create(user_id, import_type, upload):
        """Validate, extract and privately stage an uploaded document."""
        if import_type not in {"receipt", "statement"}:
            raise ValueError("Invalid document type.")

        if not upload or not upload.filename:
            raise ValueError("Choose a file to upload.")

        filename = secure_filename(upload.filename)

        ext = (
            filename.rsplit(".", 1)[-1].lower()
            if "." in filename
            else ""
        )

        if ext not in ALLOWED:
            raise ValueError(
                "Unsupported file format. Supported: CSV, XLSX, "
                "PDF and common image formats."
            )

        raw = upload.stream.read(MAX_BYTES + 1)

        if not raw:
            raise ValueError("The selected file is empty.")

        if len(raw) > MAX_BYTES:
            raise ValueError("File exceeds the 10 MB upload limit.")

        signatures = {
            "pdf": raw.startswith(b"%PDF-"),
            "png": raw.startswith(b"\x89PNG\r\n\x1a\n"),
            "jpg": raw.startswith(b"\xff\xd8\xff"),
            "jpeg": raw.startswith(b"\xff\xd8\xff"),
            "webp": (
                len(raw) > 12
                and raw[:4] == b"RIFF"
                and raw[8:12] == b"WEBP"
            ),
            "tif": raw.startswith((b"II*\x00", b"MM\x00*")),
            "tiff": raw.startswith((b"II*\x00", b"MM\x00*")),
            "xlsx": raw.startswith(b"PK"),
            "xlsm": raw.startswith(b"PK"),
        }

        if ext in signatures and not signatures[ext]:
            raise ValueError(
                "The file contents do not match the selected file type."
            )

        if ext == "csv" and (
            b"\x00" in raw
            or sum(
                byte < 9 or 13 < byte < 32
                for byte in raw[:4096]
            ) > 20
        ):
            raise ValueError(
                "The uploaded CSV does not appear to be a valid text file."
            )

        digest = hashlib.sha256(raw).hexdigest()

        existing_import = DocumentImport.query.filter_by(
            user_id=user_id,
            content_sha256=digest,
        ).first()

        if existing_import:
            raise ValueError(
                "This exact file was already uploaded. "
                "Open its existing review instead."
            )

        text = DocumentIntelligenceService._extract_text(ext, raw)

        if import_type == "statement":
            transactions = (
                DocumentIntelligenceService._statement_rows(text, ext)
            )

            # Retry OCR for scanned PDFs and PDFs whose embedded text
            # does not expose transaction rows.
            if not transactions and ext == "pdf":
                ocr_text = (
                    DocumentIntelligenceService._extract_pdf_ocr(raw)
                )

                if ocr_text and ocr_text.strip():
                    transactions = (
                        DocumentIntelligenceService._statement_rows(
                            ocr_text,
                            ext,
                        )
                    )

            if not transactions:
                raise ValueError(
                    "No transaction rows could be extracted. "
                    "For PDFs, verify that the statement is readable "
                    "and that transaction dates, descriptions and "
                    "amounts can be recognized. CSV and Excel files "
                    "require supported column headers."
                )

            if len(transactions) > MAX_TRANSACTIONS:
                raise ValueError(
                    f"A single import is limited to "
                    f"{MAX_TRANSACTIONS:,} transaction rows. "
                    "Split the statement and upload each part separately."
                )

        else:
            # ---------------------------------------------------------
            # RECEIPT OCR AND TOTAL EXTRACTION
            # ---------------------------------------------------------

            # Recognize a labelled total only when an amount follows it.
            # This helps trigger OCR when a PDF's embedded text is
            # incomplete or contains labels without usable amounts.
            amount_signal = re.compile(
                r"""
                (?<!\w)
                (?:
                    GRAND\s*TOTAL
                    | NET\s*TOTAL
                    | INVOICE\s*TOTAL
                    | TOTAL\s*(?:AMOUNT|PAID|DUE)?
                    | AMOUNT\s*(?:PAID|DUE|PAYABLE)
                    | BALANCE\s*DUE
                )
                \s*[:=\-]?\s*
                (?:₦|NGN|N|\$)?\s*
                (?:\d{1,3}(?:,\d{3})+|\d+)
                (?:\.\d{1,2})?
                |
                (?:₦|NGN|\$)\s*
                (?:\d{1,3}(?:,\d{3})+|\d+)
                (?:\.\d{1,2})?
                """,
                re.IGNORECASE | re.VERBOSE,
            )

            # Retry PDF OCR when embedded text lacks a recognizable
            # labelled total or currency-marked amount.
            if ext == "pdf" and not amount_signal.search(text or ""):
                ocr_text = (
                    DocumentIntelligenceService._extract_pdf_ocr(raw)
                )

                if ocr_text and ocr_text.strip():
                    text = ocr_text

            # Normalize whitespace for matching while retaining the
            # original extracted text for line-based item processing.
            searchable_text = re.sub(
                r"[ \t]+",
                " ",
                text or "",
            )

            lines = [
                " ".join(line.split())
                for line in (text or "").splitlines()
                if line.strip()
            ]

            # A labelled total is stronger evidence than an unlabelled
            # amount. Do not match "subtotal", "total discount",
            # "total items", "cash tendered", or "change" as the total.
            total_pattern = re.compile(
                r"""
                (?<!\w)
                (?:
                    GRAND\s*TOTAL
                    | NET\s*TOTAL
                    | INVOICE\s*TOTAL
                    | TOTAL\s*(?:AMOUNT|PAID|DUE)
                    | AMOUNT\s*(?:PAID|DUE|PAYABLE)
                    | BALANCE\s*DUE
                    | TOTAL
                )
                (?!\s*(?:ITEMS?|QTY|QUANTITY|DISCOUNT|SAVINGS|
                        CHANGE|TENDERED|RECEIVED|BEFORE\s*TAX)\b)
                \s*[:=\-]?\s*
                (?:₦|NGN|N|\$)?\s*
                (
                    (?:\d{1,3}(?:,\d{3})+|\d+)
                    (?:\.\d{1,2})?
                )
                """,
                re.IGNORECASE | re.VERBOSE,
            )

            total_match = total_pattern.search(searchable_text)

            amount = (
                _amount(total_match.group(1))
                if total_match
                else None
            )

            # Reject zero, negative, or unparseable totals.
            if amount is not None and amount <= 0:
                amount = None

            # FALLBACK:
            # If no labelled total was found, accept a Naira-marked
            # amount only when exactly one positive amount exists.
            # This intentionally avoids guessing among multiple prices.
            if amount is None:
                currency_amounts = re.findall(
                    r"₦\s*((?:\d{1,3}(?:,\d{3})+|\d+)"
                    r"(?:\.\d{1,2})?)",
                    searchable_text,
                )

                parsed_amounts = [
                    _amount(value)
                    for value in currency_amounts
                ]

                parsed_amounts = [
                    value
                    for value in parsed_amounts
                    if value is not None and value > 0
                ]

                if len(parsed_amounts) == 1:
                    amount = parsed_amounts[0]

            # Extract the date using the existing date parser.
            date_match = re.search(
                r"\b(20\d{2}-\d{1,2}-\d{1,2}"
                r"|\d{1,2}[/-]\d{1,2}[/-]\d{2,4})\b",
                searchable_text,
            )

            day = (
                _date(date_match.group(1))
                if date_match
                else None
            )

            # Preserve the existing merchant-detection behavior.
            merchant = next(
                (
                    line
                    for line in lines[:8]
                    if len(line) > 2
                    and not re.search(
                        r"receipt|invoice|tax|total|date|thank you",
                        line,
                        re.IGNORECASE,
                    )
                ),
                "Unknown merchant",
            )

            if amount is None:
                raise ValueError(
                    "Could not confidently find the receipt total. "
                    "Try a clearer image/PDF or upload a statement instead."
                )

            def labelled_amount(*labels):
                pattern = (
                    r"(?<!\w)(?:"
                    + "|".join(labels)
                    + r")\s*[:=\-]?\s*"
                    r"(?:₦|NGN|N|\$)?\s*"
                    r"((?:\d{1,3}(?:,\d{3})+|\d+)"
                    r"(?:\.\d{1,2})?)"
                )

                match = re.search(
                    pattern,
                    searchable_text,
                    re.IGNORECASE,
                )

                return (
                    _amount(match.group(1))
                    if match
                    else None
                )

            subtotal = labelled_amount(
                r"sub\s*total",
                "subtotal",
            )

            tax = labelled_amount(
                "VAT",
                "tax",
                r"sales\s+tax",
            )

            payment_match = re.search(
                r"\b(cash|debit\s+card|credit\s+card|visa|"
                r"mastercard|pos|bank\s+transfer|transfer|"
                r"mobile\s+money)\b",
                searchable_text,
                re.IGNORECASE,
            )

            items = []

            for line in lines:
                item_match = re.match(
                    r"^(.{2,100}?)\s+"
                    r"(?:x\s*)?(\d+(?:\.\d+)?)?\s*"
                    r"[₦$]?\s*([\d,]+\.\d{2})$",
                    line,
                    re.IGNORECASE,
                )

                if (
                    item_match
                    and not re.search(
                        r"total|tax|subtotal|change|balance|cash|vat",
                        item_match.group(1),
                        re.IGNORECASE,
                    )
                ):
                    line_total = _amount(item_match.group(3))

                    if line_total is not None:
                        items.append({
                            "description": (
                                item_match.group(1).strip()[:100]
                            ),
                            "quantity": float(
                                item_match.group(2) or 1
                            ),
                            "line_total": line_total,
                        })

                    if len(items) >= 100:
                        break

            receipt_details = {
                "subtotal": subtotal,
                "tax": tax,
                "total": abs(amount),
                "payment_method": (
                    payment_match.group(1).title()
                    if payment_match
                    else None
                ),
                "items": items,
            }

            row = {
                "date": day or "",
                "description": merchant[:255],
                "amount": -abs(amount),
                "category": _category(merchant, "expense"),
                "kind": "expense",
                "confirmed": False,
                "fingerprint": _fingerprint(
                    day or "",
                    -abs(amount),
                    merchant,
                ),
                "possible_duplicate": False,
                "needs_date_review": not bool(day),
                "receipt_details": receipt_details,
            }

            transactions = [row]

        # Check for possible existing transactions without importing them.
        for row in transactions:
            if not row.get("date"):
                continue

            try:
                transaction_date = date.fromisoformat(row["date"])
                amount = abs(float(row["amount"]))

                if row.get("kind") == "expense":
                    existing = Expense.query.filter_by(
                        user_id=user_id,
                        merchant=str(
                            row.get("description", "")
                        )[:150],
                        expense_date=transaction_date,
                        amount=amount,
                    ).first()
                else:
                    existing = Income.query.filter_by(
                        user_id=user_id,
                        source=str(
                            row.get("description", "")
                        )[:150],
                        received_date=transaction_date,
                        amount=amount,
                    ).first()

                if existing:
                    row["possible_duplicate"] = True

            except (ValueError, TypeError):
                row["needs_review"] = True

        # Store files outside the public static directory.
        base = Path(
            current_app.config.get("INTEGRATION_UPLOAD_DIR")
            or (
                Path(current_app.instance_path)
                / "private_imports"
            )
        ).resolve()

        static_root = (
            Path(current_app.root_path) / "static"
        ).resolve()

        if base == static_root or static_root in base.parents:
            raise RuntimeError(
                "Integration uploads must be stored outside "
                "the public static directory."
            )

        base.mkdir(
            parents=True,
            exist_ok=True,
            mode=0o700,
        )

        try:
            os.chmod(base, 0o700)
        except OSError:
            pass

        stored = base / f"{uuid.uuid4().hex}.{ext}"

        with stored.open("xb") as destination:
            destination.write(raw)

        try:
            os.chmod(stored, 0o600)
        except OSError:
            pass

        record = DocumentImport(
            user_id=user_id,
            import_type=import_type,
            original_name=filename[:255],
            stored_path=str(stored),
            content_sha256=digest,
            status="review",
            extracted_json=json.dumps(
                {
                    "transactions": transactions,
                    "source_format": ext,
                },
                ensure_ascii=False,
            ),
        )

        db.session.add(record)

        try:
            db.session.commit()

        except IntegrityError as exc:
            db.session.rollback()

            try:
                stored.unlink(missing_ok=True)
            except OSError:
                pass

            raise ValueError(
                "This exact file was uploaded already. "
                "Open the existing review instead."
            ) from exc

        except Exception:
            db.session.rollback()

            try:
                stored.unlink(missing_ok=True)
            except OSError:
                pass

            raise

        return record



    @staticmethod
    def confirm(user_id, import_id, index, data):
        """Import one user-confirmed transaction."""
        record = (
            DocumentImport.query
            .filter_by(id=import_id, user_id=user_id)
            .with_for_update()
            .first()
        )

        if not record:
            raise ValueError("Import not found.")

        payload = record.extracted()
        rows = payload.get("transactions", [])

        if index < 0 or index >= len(rows):
            raise ValueError("Transaction row not found.")

        row = rows[index]

        if row.get("confirmed"):
            raise ValueError(
                "This transaction has already been imported."
            )

        description = " ".join(
            str(
                data.get(
                    "description",
                    row.get("description", ""),
                )
            ).split()
        )[:255]

        category = " ".join(
            str(
                data.get(
                    "category",
                    row.get("category", "Other"),
                )
            ).split()
        )[:100]

        day = _date(data.get("date", row.get("date")))

        amount = _amount(
            data.get("amount", row.get("amount"))
        )

        kind = str(
            data.get("kind", row.get("kind", "expense"))
        ).lower()

        if (
            not description
            or not day
            or amount is None
            or amount == 0
            or kind not in {"income", "expense"}
        ):
            raise ValueError(
                "Date, description, positive non-zero amount, "
                "and transaction type are required."
            )
        
        # The review UI accepts positive amounts; the transaction type
        # determines the accounting direction.
        amount = abs(amount) * (
            1 if kind == "income" else -1
        )

        transaction_date = date.fromisoformat(day)

        # Find an exact existing match, but do not automatically assume
        # that two identical-looking transactions are the same event.
        duplicate = None
        duplicate_model = None

        if kind == "expense":
            duplicate = Expense.query.filter_by(
                user_id=user_id,
                merchant=description,
                expense_date=transaction_date,
                amount=abs(amount),
            ).first()
            duplicate_model = "Expense"
        else:
            duplicate = Income.query.filter_by(
                user_id=user_id,
                source=description,
                received_date=transaction_date,
                amount=abs(amount),
            ).first()
            duplicate_model = "Income"

        # Only bypass duplicate protection when the user explicitly
        # certifies that this is a separate transaction.
        certify_value = data.get("certify_not_duplicate", False)

        if isinstance(certify_value, str):
            certify_not_duplicate = (
                certify_value.strip().lower()
                in {"1", "true", "yes", "on"}
            )
        else:
            certify_not_duplicate = bool(certify_value)

        override_reason = " ".join(
            str(data.get("duplicate_override_reason", "")).split()
        )[:500]

        duplicate_override = None

        if duplicate:
            if not certify_not_duplicate:
                raise ValueError(
                    "A matching transaction already exists. "
                    "Review the existing transaction. If the bank "
                    "statement confirms this is a separate transaction, "
                    "select 'Confirm not duplicate' and provide a reason."
                )

            if not override_reason:
                raise ValueError(
                    "Provide a reason for certifying that this transaction "
                    "is not a duplicate."
                )

            from app.utils.timezone import now

            duplicate_override = {
                "certified_not_duplicate": True,
                "existing_transaction_model": duplicate_model,
                "existing_transaction_id": duplicate.id,
                "certification_reason": override_reason,
                "certified_at": now().isoformat(),
                "import_id": record.public_id,
                "statement_row_index": index,
            }

        elif certify_not_duplicate:
            # Avoid storing an override certification when no matching
            # existing transaction was found.
            duplicate_override = None


        from app.subscriptions.service import SubscriptionService

        allowed, limit_message = (
            SubscriptionService.can_add_transaction(user_id)
        )

        if not allowed:
            raise ValueError(
                limit_message
                or "Your active subscription does not allow "
                "another transaction."
            )

        if not CategoryService.is_valid_for_user(
            user_id,
            kind,
            category,
        ):
            CategoryService.ensure_system_categories()

            if not CategoryService.is_valid_for_user(
                user_id,
                kind,
                category,
            ):
                category = "Other"

        if kind == "expense":
            transaction = Expense(
                user_id=user_id,
                category=category,
                merchant=description[:150],
                description=description,
                amount=abs(amount),
                payment_method=(
                    "Receipt"
                    if record.import_type == "receipt"
                    else "Bank statement"
                ),
                expense_date=transaction_date,
                notes=(
                    f"Confirmed from {record.import_type} "
                    f"import #{record.public_id}"
                ),
            )

            db.session.add(transaction)

        else:
            transaction = Income(
                user_id=user_id,
                source=description[:150],
                category=category,
                amount=abs(amount),
                received_date=transaction_date,
                notes=(
                    f"Confirmed from statement import "
                    f"#{record.public_id}"
                ),
            )

            db.session.add(transaction)

        
        row.update({
            "date": day,
            "description": description,
            "amount": amount,
            "category": category,
            "kind": kind,
            "confirmed": True,
            "possible_duplicate": False,
            "duplicate_override": duplicate_override,
        })

        record.extracted_json = json.dumps(
            payload,
            ensure_ascii=False,
        )

        record.status = (
            "imported"
            if all(item.get("confirmed") for item in rows)
            else "partial"
        )

        db.session.commit()

        return record

    @staticmethod
    def delete(user_id, import_id):
        """Delete an import and its private stored file."""
        record = DocumentImport.query.filter_by(
            id=import_id,
            user_id=user_id,
        ).first()

        if not record:
            raise ValueError("Import not found.")

        try:
            root = Path(
                current_app.config.get("INTEGRATION_UPLOAD_DIR")
                or (
                    Path(current_app.instance_path)
                    / "private_imports"
                )
            ).resolve()

            stored = Path(record.stored_path).resolve()

            if root in stored.parents:
                stored.unlink(missing_ok=True)
            else:
                current_app.logger.warning(
                    "Refused to remove import path outside private "
                    "storage id=%s",
                    record.id,
                )

        except OSError:
            current_app.logger.warning(
                "Could not remove private import file id=%s",
                record.id,
            )

        db.session.delete(record)
        db.session.commit()